import streamlit as st
import pandas as pd
import openpyxl
import io
import math
import datetime

# --- CONFIGURACIÓN DE PÁGINA ---
st.set_page_config(
    page_title="Generador de Bitácoras de Combustible",
    page_icon="⛽",
    layout="wide"
)

st.title("⛽ Generador Automatizado de Bitácoras de Combustible")
st.caption("Dirección de Desarrollo Urbano y Medio Ambiente | H. Ayuntamiento de Campeche")

# --- PARÁMETROS Y CATÁLOGOS BASE ---
VEHICULOS_DEFAULT = [
    {"direccion": "MEDIO AMBIENTE", "placa": "DKR456B", "tipo": "JIMNY GLX TA, 5 PUERTAS 2025", "combustible": "PREMIUM", "importe_total": 5000.00, "tope_diario": 800.0, "anio": "2025", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "DFT-565-C", "tipo": "JETTA", "combustible": "MAGNA", "importe_total": 5300.00, "tope_diario": 1000.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "CN-2633-B", "tipo": "CAMIONETA (RAM)", "combustible": "MAGNA", "importe_total": 7049.27, "tope_diario": 2000.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "88GWU7", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 1112.75, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "87GWU8", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 800.00, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "88GWU8", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 750.00, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "90GWU8", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 700.00, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "86GWU8", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 688.13, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "86GWU7", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 683.46, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "89GWU7", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 500.00, "tope_diario": 200.0, "anio": "", "activo": True},
    {"direccion": "MEDIO AMBIENTE", "placa": "89GWU8", "tipo": "MOTOCICLETA", "combustible": "MAGNA", "importe_total": 500.00, "tope_diario": 200.0, "anio": "", "activo": True},
]

PRECIOS = {
    "MAGNA": 23.99,
    "PREMIUM": 26.95,
    "DIESEL": 27.00
}

# --- BARRA LATERAL: PARÁMETROS DEL PERIODO ---
with st.sidebar:
    st.header("⚙️ Configuración del Periodo")
    fecha_ini = st.date_input("Inicio del periodo", datetime.date(2026, 9, 26))
    fecha_fin = st.date_input("Fin del periodo", datetime.date(2026, 10, 25))
    
    fecha_entrega_txt = st.text_input("Texto de Entrega (Fila 27)", "25 DE OCTUBRE DE 2026")
    destino_txt = st.text_input("Destino o Uso (Fila 28)", "RECORRIDOS OPERATIVOS, SUPERVISIÓN E INSPECCIÓN EN EL MUNICIPIO")
    responsable_txt = st.text_input("Responsable Operativo sugerido", "OPERADOR ASIGNADO")
    
    st.markdown("---")
    st.subheader("📁 Plantilla Excel")
    uploaded_file = st.file_uploader("Subir '1.FORMATO BITACORA 2024 - 2027.xlsx' (opcional si ya está en la raíz)", type=["xlsx"])

# Días autorizados: Lunes(0), Martes(1), Jueves(3), Viernes(4)
dias_autorizados = [0, 1, 3, 4]
fechas_disponibles = []
cur = fecha_ini
while cur <= fecha_fin:
    if cur.weekday() in dias_autorizados:
        fechas_disponibles.append(cur)
    cur += datetime.timedelta(days=1)

with st.sidebar:
    st.info(f"📅 **{len(fechas_disponibles)} fechas válidas** encontradas (Lunes, Martes, Jueves y Viernes).")

# --- TABLA EDITABLE ---
st.subheader("1. Cuotas Asignadas y Límites Diarios")
df_edit = st.data_editor(
    pd.DataFrame(VEHICULOS_DEFAULT),
    column_config={
        "activo": st.column_config.CheckboxColumn("Generar", default=True),
        "importe_total": st.column_config.NumberColumn("Presupuesto ($)", format="$%.2f"),
        "tope_diario": st.column_config.NumberColumn("Tope Día ($)", format="$%.2f"),
        "combustible": st.column_config.SelectboxColumn("Combustible", options=["MAGNA", "PREMIUM", "DIESEL"])
    },
    use_container_width=True,
    num_rows="dynamic"
)

# Resumen numérico
tot_magna = df_edit[df_edit["activo"] & (df_edit["combustible"] == "MAGNA")]["importe_total"].sum()
tot_prem = df_edit[df_edit["activo"] & (df_edit["combustible"] == "PREMIUM")]["importe_total"].sum()
col1, col2, col3 = st.columns(3)
col1.metric("Total Magna (CD)", f"${tot_magna:,.2f}")
col2.metric("Total Premium (Vales)", f"${tot_prem:,.2f}")
col3.metric("Total General", f"${tot_magna + tot_prem:,.2f}")

# --- LÓGICA DE DISTRIBUCIÓN ---
def calcular_cargas(total_monto, tope_dia, fechas, combustible):
    precio_unit = PRECIOS.get(combustible, 23.99)
    n_cargas = math.ceil(total_monto / tope_dia)
    
    # Máximo 11 renglones disponibles en la bitácora (filas 14 a 24)
    if n_cargas > 11:
        n_cargas = 11
        tope_dia = math.ceil(total_monto / 11)

    step = len(fechas) / n_cargas
    fechas_asignadas = [fechas[int(round(i * step))] for i in range(n_cargas)]
    
    cargas = []
    restante = round(total_monto, 2)
    
    for i in range(n_cargas):
        if i == n_cargas - 1:
            importe = round(restante, 2)
        else:
            importe = min(tope_dia, round(restante - (n_cargas - 1 - i) * 1.0, 2))
            if importe > tope_dia:
                importe = tope_dia
        
        litros = round(importe / precio_unit, 2)
        cargas.append({
            "fecha": fechas_asignadas[i].strftime("%d/%m/%Y"),
            "litros": litros,
            "precio_unit": precio_unit,
            "importe": importe
        })
        restante = round(restante - importe, 2)
        
    return cargas

def escribir_hoja_bitacora(ws, datos, cargas):
    # Encabezado (Fila 9)
    ws["B9"] = f"N° ECONOMICO: {datos.get('num_eco', '')}"
    ws["E9"] = f"MARCA Y TIPO: {datos.get('tipo', '')}"
    ws["I9"] = f"AÑO: {datos.get('anio', '')}"
    ws["L9"] = f"PLACAS: {datos.get('placa', '')}"
    
    # Renglones de recargas (14 a 24)
    for idx, c in enumerate(cargas):
        r = 14 + idx
        if r > 24:
            break
        ws[f"B{r}"] = c["fecha"]
        ws[f"C{r}"] = c["litros"]
        ws[f"E{r}"] = c["precio_unit"]
        ws[f"F{r}"] = c["importe"]
        ws[f"J{r}"] = responsable_txt
        
    # Fila 25: Fórmulas nativas de Excel
    ws["C25"] = "=SUM(C14:C24)"
    ws["F25"] = "=SUM(F14:F24)"
    
    # Fila 27 y 28: Pie de entrega y justificación
    ws["B27"] = f"BITÁCORA CORRESPONDIENTE A LA ENTREGA DEL DÍA: {fecha_entrega_txt}"
    ws["B28"] = f"DESTINO O USO DEL VEHICULO: {destino_txt}"

# --- GENERADOR DE EXCEL ---
st.markdown("---")
st.subheader("2. Procesamiento del Libro Excel")

if st.button("🚀 Generar Bitácoras Completas"):
    # Carga de plantilla base
    if uploaded_file is not None:
        wb = openpyxl.load_workbook(uploaded_file)
    else:
        try:
            wb = openpyxl.load_workbook("1.FORMATO BITACORA 2024 - 2027.xlsx")
        except FileNotFoundError:
            st.error("No se encontró el archivo '1.FORMATO BITACORA 2024 - 2027.xlsx' en la raíz. Por favor súbelo en la barra lateral izquierda.")
            st.stop()
            
    ws_base = wb["BITACORA VEHICULOS"]
    
    # Procesar cada vehículo
    for _, row in df_edit.iterrows():
        if not row["activo"]:
            continue
            
        cargas = calcular_cargas(
            row["importe_total"],
            row["tope_diario"],
            fechas_disponibles,
            row["combustible"]
        )
        
        # Clonar la pestaña manteniendo diseño y firmas intactas
        ws_nueva = wb.copy_worksheet(ws_base)
        ws_nueva.title = str(row["placa"]).replace("-", "").strip()[:31]
        escribir_hoja_bitacora(ws_nueva, row.to_dict(), cargas)
        
    # Guardar en buffer en memoria
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    st.success("¡Bitácoras generadas correctamente para todos los vehículos seleccionados!")
    st.download_button(
        label="📥 Descargar Libro Excel con Bitácoras Justificadas",
        data=buffer,
        file_name="BITACORAS_JUSTIFICADAS_OCTUBRE_2026.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
